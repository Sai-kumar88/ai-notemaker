"""
app/services/document_service.py

High-performance text extraction engine for PDF and DOCX textbooks and documents.
Uses PyMuPDF (fitz) as primary engine with graceful fallback to pypdf for PDFs,
and python-docx for Word documents.
"""
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional

import docx
import pypdf

from app.config import settings

logger = logging.getLogger(__name__)


class DocumentExtractionError(Exception):
    """Raised when text cannot be extracted from the document."""
    pass


class UnsupportedFileTypeError(Exception):
    """Raised when uploaded file type is not supported."""
    pass


@dataclass
class DocumentPage:
    """Represents an individual page within an extracted document."""
    page_number: int
    text: str


@dataclass
class ExtractedDocument:
    """Represents a fully extracted document with metadata and page segments."""
    filename: str
    file_type: str
    full_text: str
    pages: List[DocumentPage]
    word_count: int = field(init=False)
    char_count: int = field(init=False)

    def __post_init__(self):
        self.word_count = len(self.full_text.split())
        self.char_count = len(self.full_text)


class DocumentService:
    """
    Orchestrates validation and text extraction across supported document formats.
    """

    @staticmethod
    def validate_file(filename: str, file_size: int) -> str:
        """
        Validates file extension and size against environment settings.
        No upload limits or allowed extensions are hardcoded in business logic.
        """
        ext = Path(filename).suffix.lower()
        if ext not in settings.allowed_extensions_set:
            allowed_str = ", ".join(settings.allowed_extensions)
            raise UnsupportedFileTypeError(
                f"Unsupported file format '{ext}'. Allowed formats: {allowed_str}"
            )

        if file_size > settings.max_file_size_bytes:
            raise DocumentExtractionError(
                f"File size ({file_size / (1024 * 1024):.2f} MB) exceeds maximum allowed size "
                f"({settings.max_file_size_mb} MB)."
            )

        return ext

    @classmethod
    def extract_text_from_file(cls, file_path: Path, original_filename: str) -> ExtractedDocument:
        """
        Extracts structured text and pages from a PDF or DOCX file.
        """
        ext = cls.validate_file(original_filename, file_path.stat().st_size)

        if ext == ".pdf":
            return cls._extract_from_pdf(file_path, original_filename)
        elif ext == ".docx":
            return cls._extract_from_docx(file_path, original_filename)
        else:
            raise UnsupportedFileTypeError(f"Unsupported file type: {ext}")

    @staticmethod
    def _extract_from_pdf(file_path: Path, filename: str) -> ExtractedDocument:
        """
        Extracts text from PDF.
        Priority:
          1. PyMuPDF (fitz) - Industry gold standard for textbooks & complex layouts
          2. pypdf - Robust pure-Python fallback
        """
        pages: List[DocumentPage] = []
        full_text_parts: List[str] = []

        # Engine 1: PyMuPDF (fitz)
        try:
            import fitz
            doc = fitz.open(str(file_path))
            try:
                if len(doc) == 0:
                    raise DocumentExtractionError("PDF file contains no pages.")

                for idx, page in enumerate(doc, start=1):
                    try:
                        page_text = page.get_text("text") or ""
                        page_text = page_text.strip()
                        if page_text:
                            pages.append(DocumentPage(page_number=idx, text=page_text))
                            full_text_parts.append(page_text)
                    except Exception as p_err:
                        logger.warning("PyMuPDF failed on page %d of %s: %s", idx, filename, p_err)
            finally:
                doc.close()
        except DocumentExtractionError:
            raise
        except Exception as fitz_err:
            logger.warning("PyMuPDF extraction issue on %s (%s), attempting fallback...", filename, fitz_err)

        # Engine 2: Fallback to pypdf if PyMuPDF extracted no text
        if not full_text_parts:
            try:
                reader = pypdf.PdfReader(str(file_path), strict=False)
                if len(reader.pages) == 0:
                    raise DocumentExtractionError("PDF file contains no pages.")

                for idx, page in enumerate(reader.pages, start=1):
                    try:
                        page_text = page.extract_text() or ""
                        page_text = page_text.strip()
                        if page_text:
                            pages.append(DocumentPage(page_number=idx, text=page_text))
                            full_text_parts.append(page_text)
                    except Exception as page_err:
                        logger.warning("pypdf error on page %d of %s: %s", idx, filename, page_err)
            except DocumentExtractionError:
                raise
            except Exception as pypdf_err:
                raise DocumentExtractionError(f"Failed to extract text from PDF: {str(pypdf_err)}") from pypdf_err

        full_text = "\n\n".join(full_text_parts).strip()
        if not full_text:
            raise DocumentExtractionError(
                "No readable text could be extracted from the PDF. "
                "The document might be image-only (scanned), empty, or password-protected."
            )

        return ExtractedDocument(
            filename=filename,
            file_type="pdf",
            full_text=full_text,
            pages=pages
        )

    @staticmethod
    def _extract_from_docx(file_path: Path, filename: str) -> ExtractedDocument:
        """
        Extracts text and table contents from DOCX files.
        """
        try:
            doc = docx.Document(str(file_path))
            paragraphs: List[str] = []

            for p in doc.paragraphs:
                text = p.text.strip()
                if text:
                    paragraphs.append(text)

            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        # De-duplicate consecutive identical cells from merged columns
                        deduped = []
                        for c in row_cells:
                            if not deduped or deduped[-1] != c:
                                deduped.append(c)
                        row_text = " | ".join(deduped)
                        if row_text:
                            paragraphs.append(row_text)

            full_text = "\n\n".join(paragraphs).strip()
            if not full_text:
                raise DocumentExtractionError("DOCX file does not contain readable text.")

            pages = [DocumentPage(page_number=1, text=full_text)]

            return ExtractedDocument(
                filename=filename,
                file_type="docx",
                full_text=full_text,
                pages=pages
            )
        except DocumentExtractionError:
            raise
        except Exception as e:
            raise DocumentExtractionError(f"Failed to extract text from DOCX: {str(e)}") from e
